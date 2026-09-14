"""Build the team contribution record for the CMPE 272 submission."""

# Author: Mathew — Word document generation for the team contribution record.

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path("Team_Contributions.docx")
NAVY = "1F4E78"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"


def set_cell_shading(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def remove_paragraph_borders(style) -> None:
    paragraph_properties = style.element.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def set_cell_border(cell, color: str = LIGHT_GRAY) -> None:
    tc_borders = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        cell._tc.get_or_add_tcPr().append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def style_run(run, *, bold: bool = False, size: int = 11, color: str = "000000") -> None:
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def write_cell(cell, text: str, *, header: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_run(paragraph.add_run(text), bold=header, size=10, color="FFFFFF" if header else "000000")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell, NAVY if header else LIGHT_BLUE)
    set_cell_border(cell)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    style_run(run, bold=True, size=14)


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    remove_paragraph_borders(document.styles["Title"])

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("GitHub Issues Service Team Contributions")
    style_run(title_run, bold=True, size=22)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    style_run(subtitle.add_run("CMPE 272 Assignment"), size=11)

    introduction = document.add_paragraph()
    introduction.paragraph_format.space_after = Pt(10)
    style_run(
        introduction.add_run(
            "This document records the team division of work for the GitHub Issues Service. "
            "The assignments below match the authorship comments in the source files."
        )
    )

    add_heading(document, "Individual Contributions")
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Inches(1.0)
    table.columns[1].width = Inches(3.65)
    table.columns[2].width = Inches(2.35)
    headers = ("Member", "Work completed", "Primary code and artifacts")
    for cell, text in zip(table.rows[0].cells, headers):
        write_cell(cell, text, header=True)

    rows = (
        (
            "Sonit",
            "Set up the FastAPI application, environment configuration, request and response models, "
            "health endpoint, request IDs, and shared error responses.",
            "app/config.py\napp/models.py\napp/main.py",
        ),
        (
            "Sagar",
            "Implemented the GitHub REST client, issue and comment endpoints, pagination forwarding, "
            "and GitHub error and rate limit translation.",
            "app/github.py\napp/issues.py",
        ),
        (
            "Joel",
            "Implemented GitHub webhook signature verification, supported event and action validation, "
            "delivery deduplication, event storage, and webhook routes.",
            "app/webhooks.py\napp/webhook_routes.py",
        ),
        (
            "Mathew",
            "Wrote unit and live integration tests, prepared the OpenAPI contract, Docker and Make setup, "
            "README, design note, and this contribution record.",
            "tests/\nopenapi.yaml\nDockerfile\nREADME.md\nDESIGN.md",
        ),
    )
    for member, work, artifacts in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (member, work, artifacts)):
            write_cell(cell, text)

    for row in table.rows:
        row.cells[0].width = Inches(1.0)
        row.cells[1].width = Inches(3.65)
        row.cells[2].width = Inches(2.35)

    add_heading(document, "Source Code Attribution")
    attribution = document.add_paragraph()
    attribution.paragraph_format.space_after = Pt(8)
    style_run(
        attribution.add_run(
            "Each application module and test file begins with a # Author comment. "
            "These comments identify the team member responsible for that file or group of related work."
        )
    )

    add_heading(document, "Shared Review")
    review = document.add_paragraph()
    review.paragraph_format.space_after = Pt(0)
    style_run(
        review.add_run(
            "The team reviewed the API contract, test results, webhook behavior, and documentation "
            "together before submission."
        )
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    main()
